"""Tool 4 — XLSX writer.

Fill the rule-set's output spreadsheet template from the check results. The
template already contains the rules and gates as rows; this tool only writes
the verdict (PASSED/FAILED/ERROR) into each rule's `Result` cell.

Template contract (row 1 is the header):
  - a `Rule` column (`rule`, `rule name`, or `name`) is the row key, matched
    against `CheckResult.rule_name`;
  - a `Result` column (`result`, `verdict`, `status`, or `outcome`) is the cell
    that receives the verdict.
"""

from __future__ import annotations

import logging
from pathlib import Path

import yaml
from openpyxl import load_workbook

from .config import Config
from .schemas import CheckResult

logger = logging.getLogger(__name__)

_RULE_HEADERS = {"rule", "rule name", "name"}
_RESULT_HEADERS = {"result", "verdict", "status", "outcome"}

_VERDICT_LABELS = {"pass": "PASSED", "fail": "FAILED", "error": "ERROR"}


def load_output_template_names(config: Config) -> list[str]:
    """Return the template filenames from the `output-template` YAML section."""
    data = yaml.safe_load(config.rules_path.read_text(encoding="utf-8")) or {}
    names: list[str] = []
    for entry in data.get("output-template", []):
        name = entry.get("file") if isinstance(entry, dict) else entry
        if name:
            names.append(name)
    return names


def _find_column(headers: list, candidates: set[str]) -> int | None:
    for index, header in enumerate(headers):
        if str(header).strip().lower() in candidates:
            return index
    return None


def _fill_template(template_path: Path, out_path: Path, checks: list[CheckResult]) -> None:
    workbook = load_workbook(template_path)
    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]
    rule_col = _find_column(headers, _RULE_HEADERS)
    result_col = _find_column(headers, _RESULT_HEADERS)
    if rule_col is None or result_col is None:
        logger.warning(
            "Template %s is missing Rule/Result header columns; not filling",
            template_path,
        )
        return

    results_by_rule = {check.rule_name: check for check in checks}
    unmatched: list[str] = []
    for row in sheet.iter_rows(min_row=2):
        rule_cell = row[rule_col] if rule_col < len(row) else None
        if rule_cell is None or rule_cell.value is None:
            continue
        rule_name = str(rule_cell.value).strip()
        if not rule_name:
            continue
        check = results_by_rule.get(rule_name)
        if check is None:
            unmatched.append(rule_name)
            continue
        result_cell = row[result_col] if result_col < len(row) else None
        if result_cell is None:
            logger.warning("No Result cell in row for rule '%s'", rule_name)
            continue
        result_cell.value = _VERDICT_LABELS.get(check.verdict, "ERROR")

    if unmatched:
        logger.warning(
            "No check result for template rule(s): %s", ", ".join(unmatched)
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def write_output_xlsx(
    config: Config, template_names: list[str], checks: list[CheckResult]
) -> list[str]:
    """Fill each listed template; write `output.xlsx`, `output-2.xlsx`, ...

    Returns the list of output filenames written into the work dir.
    """
    outputs: list[str] = []
    for index, name in enumerate(template_names):
        template_path = config.skills_dir / name
        if not template_path.exists():
            logger.warning("Output template %s not found; skipping", template_path)
            continue
        out_name = "output.xlsx" if not outputs else f"output-{index + 1}.xlsx"
        out_path = config.work_dir / out_name
        logger.info("Filling output template %s -> %s", template_path, out_path)
        _fill_template(template_path, out_path, checks)
        outputs.append(out_name)
    return outputs
